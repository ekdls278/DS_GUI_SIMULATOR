"""Local-only A:E Excel ICD import. No workbook contents are uploaded or evaluated."""
from dataclasses import asdict, dataclass
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException


@dataclass(frozen=True)
class IcdField:
    key: str
    sequence: str
    name: str
    choices: tuple[str, ...]
    default: str

    @property
    def label(self):
        return f'{self.sequence} · {self.name}' if self.sequence else self.name


@dataclass(frozen=True)
class IcdMessage:
    opcode: str
    description: str
    elements: tuple[IcdField, ...]
    src: str = ''
    dst: str = ''

    @property
    def cmd(self):
        return self.opcode

    @property
    def fields(self):
        return tuple((f.key, '선택' if f.choices else '텍스트', f.default, '—', f.label) for f in self.elements)

    def field_rows(self):
        return [[f.sequence, f.name, '선택' if f.choices else '직접 입력', ' / '.join(f.choices), f.default] for f in self.elements]


def message_id(opcode, name):
    return 'icd_' + hashlib.sha256((opcode + '\0' + name).encode('utf-8')).hexdigest()


def _book(path):
    path = Path(path)
    if path.suffix.lower() not in ('.xlsx', '.xlsm'):
        raise ValueError('Excel .xlsx 또는 .xlsm 파일을 선택하세요. .xls는 .xlsx로 저장해 사용하세요.')
    try:
        return load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, KeyError, ValueError, ParseError, DefusedXmlException) as exc:
        raise ValueError('Excel 파일을 읽을 수 없습니다. 파일 형식을 확인하세요.') from exc


def data_sheet_names(path):
    book = _book(path)
    try:
        return [sheet.title for sheet in book.worksheets if sheet.title.rstrip().endswith('(데이터)')]
    finally:
        book.close()


def _cell_text(cell):
    if cell.value is None:
        return ''
    if cell.data_type == 'f':
        raise ValueError(f'{cell.coordinate}: 수식 대신 계산된 값을 사용하세요. 프로그램에서 수식을 실행하지 않습니다.')
    value = cell.value
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        # Preserve explicit Excel zero padding without guessing decimal vs hexadecimal.
        if isinstance(value, int) and re.fullmatch(r'0+', cell.number_format or ''):
            return f'{value:0{len(cell.number_format)}d}'
    return str(value).strip()


def read_icd_sheet(path, sheet_name):
    book = _book(path)
    try:
        if not isinstance(sheet_name, str) or sheet_name not in {s.title for s in book.worksheets} or not sheet_name.rstrip().endswith('(데이터)'):
            raise ValueError('(데이터)로 끝나는 시트 하나를 선택하세요.')
        sheet = book[sheet_name]
        sheet.reset_dimensions()
        groups = {}
        nonempty_rows = 0
        opcode = name = ''
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=5), 2):
            a, b, sequence, element, details = [_cell_text(cell) for cell in cells]
            if not any((a, b, sequence, element, details)):
                continue
            nonempty_rows += 1
            if nonempty_rows > 50000:
                raise ValueError('한 시트는 최대 50,000개 데이터 행을 지원합니다.')
            if a:
                if a != opcode:
                    name = ''
                opcode = a
            if b:
                name = b
            if not opcode or not name:
                raise ValueError(f'{row_number}행: opcode와 연동자료명이 필요합니다. 첫 행의 A·B열을 확인하세요.')
            key = message_id(opcode, name)
            if key not in groups:
                groups[key] = [opcode, name, []]
            if not element:
                if sequence or details:
                    raise ValueError(f'{row_number}행: D열 데이터 요소가 없습니다.')
                continue
            choices = tuple(dict.fromkeys(details.split()))
            if len(choices) > 1000:
                raise ValueError(f'{row_number}행: 선택값은 최대 1,000개입니다.')
            fields = groups[key][2]
            fields.append(IcdField(f'field_{len(fields) + 1:04}', sequence, element, choices, choices[0] if choices else ''))
        if not groups:
            raise ValueError('선택한 시트의 2행 이후 A~E열에 ICD 데이터가 없습니다.')
        if len(groups) > 10000:
            raise ValueError('한 시트는 최대 10,000개 메시지를 지원합니다.')
        return {key: IcdMessage(opcode, name, tuple(fields)) for key, (opcode, name, fields) in groups.items()}
    except (BadZipFile, ParseError, DefusedXmlException) as exc:
        raise ValueError('선택한 시트의 데이터를 읽을 수 없습니다. Excel 파일을 확인하세요.') from exc
    finally:
        book.close()


def catalog_payload(messages):
    rows = []
    for key, message in messages.items():
        if not isinstance(message, IcdMessage):
            raise ValueError('가져온 ICD 정의만 저장할 수 있습니다.')
        rows.append({'id': key, 'opcode': message.opcode, 'name': message.description, 'fields': [{**asdict(field), 'choices': list(field.choices)} for field in message.elements]})
    canonical = json.dumps(rows, sort_keys=True, ensure_ascii=False, separators=(',', ':'))
    return {'format': 'ds-icd-catalog', 'version': 1, 'fingerprint': hashlib.sha256(canonical.encode('utf-8')).hexdigest(), 'messages': rows}


def catalog_from_payload(payload):
    if not isinstance(payload, dict) or payload.get('format') != 'ds-icd-catalog' or type(payload.get('version')) is not int or payload['version'] != 1:
        raise ValueError('지원하지 않는 ICD 정의 형식입니다.')
    rows = payload.get('messages')
    if not isinstance(rows, list) or not 1 <= len(rows) <= 10000:
        raise ValueError('ICD 메시지 구성이 올바르지 않습니다.')
    result = {}
    total = 0
    try:
        for row in rows:
            opcode, name, fields = row['opcode'], row['name'], row['fields']
            if not isinstance(opcode, str) or not opcode or not isinstance(name, str) or not name or not isinstance(fields, list):
                raise ValueError('opcode·연동자료명·데이터 요소 구성을 확인하세요.')
            key = message_id(opcode, name)
            if row['id'] != key or key in result:
                raise ValueError('ICD ID가 일치하지 않거나 중복됩니다.')
            elements = []
            seen = set()
            for field in fields:
                total += 1
                if total > 50000 or not all(isinstance(field[k], str) for k in ('key', 'sequence', 'name', 'default')):
                    raise ValueError('ICD 필드 구성이 올바르지 않습니다.')
                choices = field['choices']
                if not isinstance(choices, list) or len(choices) > 1000 or not all(isinstance(v, str) and v for v in choices) or len(set(choices)) != len(choices):
                    raise ValueError('ICD 선택값 구성이 올바르지 않습니다.')
                if not field['key'] or field['key'] in seen or not field['name'] or (choices and field['default'] not in choices):
                    raise ValueError('ICD 필드 ID 또는 초기값을 확인하세요.')
                seen.add(field['key'])
                elements.append(IcdField(field['key'], field['sequence'], field['name'], tuple(choices), field['default']))
            result[key] = IcdMessage(opcode, name, tuple(elements))
    except (KeyError, TypeError, AttributeError):
        raise ValueError('ICD 정의 파일 구조가 올바르지 않습니다.') from None
    if catalog_payload(result)['fingerprint'] != payload.get('fingerprint'):
        raise ValueError('저장된 ICD 정의의 내용과 식별값이 일치하지 않습니다.')
    return result
