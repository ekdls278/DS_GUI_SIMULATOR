import csv
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from PySide6.QtCore import Qt, QTimer
from PySide6.QtGui import QFont
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication

from ds_simulator.app import SimulatorWindow
from ds_simulator.domain import Scenario, Step
from ds_simulator.preview_data import CONNECTIONS, SOCKETS, ConnectionPreview
from ds_simulator.runtime import MockTransport, ScenarioRunner
from ds_simulator.scenario_editor import ROLE
from ds_simulator.theme import STYLESHEET

APP = QApplication.instance() or QApplication([])
APP.setStyle('Fusion')
APP.setFont(QFont('Noto Sans CJK KR', 10))
APP.setStyleSheet(STYLESHEET)


def wait_until(predicate, timeout=1500):
    for _ in range(timeout // 10):
        if predicate():
            return
        QTest.qWait(10)
    raise AssertionError('Timed out waiting for Qt operation')


class RunnerTests(unittest.TestCase):
    def setUp(self):
        self.sockets = list(SOCKETS)
        self.transport = MockTransport(lambda: list(CONNECTIONS), lambda: self.sockets)
        self.events = []
        self.transport.sent.connect(self.events.append)
        self.runner = ScenarioRunner(self.transport)

    def tearDown(self):
        self.runner.stop()
        self.runner.deleteLater()
        self.transport.deleteLater()
        APP.processEvents()

    def test_three_repeats_share_sender_and_preserve_order(self):
        scenario = Scenario.example()
        scenario.steps[0].count = 3
        scenario.steps[0].children[1].milliseconds = 1
        self.runner.start(scenario)
        wait_until(lambda: not self.runner.running)
        self.assertEqual(self.runner.completed, 9)
        self.assertEqual([e['message'] for e in self.events], ['SET_FREQUENCY', 'REQUEST_STATUS'] * 3)
        self.assertTrue(all(e['mode'] == 'MOCK' and e['source'] == '시나리오' for e in self.events))

    def test_stop_cancels_pending_wait_and_future_sends(self):
        self.runner.start(Scenario(steps=[Step('wait', milliseconds=1000), Step.send()]))
        QTest.qWait(20)
        self.runner.stop()
        QTest.qWait(40)
        self.assertFalse(self.runner.timer.isActive())
        self.assertEqual(self.runner.completed, 0)
        self.assertEqual(self.events, [])

    def test_disconnection_during_wait_stops_without_sending_elsewhere(self):
        states = []
        self.runner.finished.connect(states.append)
        self.runner.start(Scenario(steps=[Step('wait', milliseconds=30), Step.send()]))
        self.sockets = [s for s in self.sockets if s.connection_id != 'TCP-01']
        wait_until(lambda: not self.runner.running)
        self.assertTrue(states[-1].startswith('오류:'))
        self.assertEqual(self.events, [])

    def test_server_listener_cannot_be_used_for_transmission(self):
        step = Step.send()
        step.connection_id = 'SRV-01'
        for identifier in ('', 'SOCK-003'):
            step.socket_id = identifier
            with self.assertRaises(ValueError):
                self.transport.send(step)
        step.socket_id = 'SOCK-005'
        event = self.transport.send(step)
        self.assertEqual(event['peer'], '192.168.0.22:51002')

    def test_snapshot_is_independent_of_editor_changes(self):
        step = Step.send()
        scenario = Scenario(steps=[Step('wait', milliseconds=10), step])
        self.runner.start(scenario)
        step.values['channel'] = '9'
        wait_until(lambda: not self.runner.running)
        self.assertEqual(self.events[0]['values']['channel'], 1)

    def test_empty_repeat_is_not_executed(self):
        with self.assertRaises(ValueError):
            self.runner.start(Scenario(steps=[Step('repeat', children=[]), Step.send()]))
        self.assertFalse(self.runner.running)


class GuiTests(unittest.TestCase):
    def setUp(self):
        self.window = SimulatorWindow(mock=True)
        self.window.show()
        APP.processEvents()

    def tearDown(self):
        self.window.scenario_editor.dirty = False
        self.window.close()
        self.window.deleteLater()
        APP.processEvents()

    def connect_first(self):
        QTest.mouseClick(self.window.connect_button, Qt.MouseButton.LeftButton)
        self.window.message_catalog.setCurrentRow(0)
        APP.processEvents()
        self.assertTrue(self.window.send_button.isEnabled())

    def click_peer(self, card, index):
        item = card.root.child(index)
        card.root.setExpanded(True)
        card.tree.scrollToItem(item)
        APP.processEvents()
        QTest.mouseClick(card.tree.viewport(), Qt.MouseButton.LeftButton, pos=card.tree.visualItemRect(item).center())
        APP.processEvents()

    def test_server_tree_selects_exact_client_and_header_is_not_send_target(self):
        self.window.select_connection(2)
        card = self.window.connection_cards['SRV-01']
        self.assertTrue(card.is_server)
        self.assertIn('0.0.0.0:7000', card.local_value.text())
        self.assertEqual(card.root.childCount(), 2)
        self.assertIn('192.168.0.21:51001', card.root.child(0).text(0))
        self.assertIn('192.168.0.22:51002', card.root.child(1).text(0))
        self.assertFalse(self.window.send_button.isEnabled())
        self.click_peer(card, 1)
        self.assertEqual(self.window.current_socket().identifier, 'SOCK-005')
        self.window.message_catalog.setCurrentRow(0)
        QTest.mouseClick(self.window.send_button, Qt.MouseButton.LeftButton)
        self.assertEqual(self.window.log_events[0]['peer'], '192.168.0.22:51002')
        self.click_peer(card, 0)
        self.assertIsNone(self.window.selected_message_name)
        self.assertFalse(self.window.send_button.isEnabled())
        self.assertEqual(self.window.current_socket().identifier, 'SOCK-004')
        QTest.mouseClick(card.header, Qt.MouseButton.LeftButton)
        self.assertIsNone(self.window.current_socket())
        self.assertEqual(card.tree.selectedItems(), [])

    def test_multiple_servers_have_separate_children(self):
        self.window.add_connection_preview(ConnectionPreview('SRV-02', '두 번째 서버', 'TCP Server', '127.0.0.1:7100'))
        second = self.window.connection_cards['SRV-02']
        first = self.window.connection_cards['SRV-01']
        self.assertEqual(second.root.childCount(), 0)
        self.window.toggle_preview_connection()
        self.window.accept_preview_peer()
        self.assertEqual(second.root.childCount(), 1)
        self.assertEqual(first.root.childCount(), 2)
        self.assertEqual(self.window.current_socket().connection_id, 'SRV-02')
        self.click_peer(first, 0)
        self.assertEqual(self.window.selected_connection().identifier, 'SRV-01')
        self.assertEqual(self.window.current_socket().identifier, 'SOCK-004')
        self.assertEqual(second.tree.selectedItems(), [])

    def test_server_stop_and_restart_remove_stale_children_and_selection(self):
        self.window.select_connection(2)
        card = self.window.connection_cards['SRV-01']
        self.click_peer(card, 0)
        self.window.message_catalog.setCurrentRow(0)
        self.window.toggle_preview_connection()
        self.assertEqual(card.root.childCount(), 0)
        self.assertIn('중지됨', card.status_badge.text())
        self.assertFalse(self.window.send_button.isEnabled())
        self.window.toggle_preview_connection()
        self.assertIn('대기중', card.status_badge.text())
        self.assertIsNone(self.window.current_socket())
        self.window.accept_preview_peer()
        self.assertEqual(card.root.childCount(), 1)
        self.assertEqual(self.window.current_socket().identifier, card.root.child(0).data(0, Qt.ItemDataRole.UserRole))

    def test_disappearing_peer_clears_tree_and_composer_target(self):
        self.window.select_connection(2)
        card = self.window.connection_cards['SRV-01']
        self.click_peer(card, 0)
        self.window.message_catalog.setCurrentRow(0)
        self.window.sockets = [s for s in self.window.sockets if s.identifier != 'SOCK-004']
        self.window.refresh_socket_monitor()
        self.assertEqual(card.root.childCount(), 1)
        self.assertIsNone(self.window.current_socket())
        self.assertIsNone(self.window.selected_message_name)
        self.assertFalse(self.window.send_button.isEnabled())

    def test_profile_reload_replaces_server_card_groups_without_duplicate_widgets(self):
        old_card = self.window.connection_cards['SRV-01']
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'connections.json'
            self.window.save_profiles_to(path)
            self.window.load_profiles_from(path)
        self.assertIsNot(self.window.connection_cards['SRV-01'], old_card)
        self.assertEqual(len(self.window.connection_cards), 4)
        self.assertEqual(self.window.connection_list.count(), 5)  # four cards and stretch
        self.assertEqual(self.window.connection_cards['SRV-01'].root.childCount(), 0)

    def test_spectrum_popup_is_nonmodal_singleton_and_does_not_navigate(self):
        current = self.window.stack.currentIndex()
        QTest.mouseClick(self.window.nav_group.button(3), Qt.MouseButton.LeftButton)
        popup = self.window.spectrum_window
        self.assertTrue(popup.isVisible())
        self.assertFalse(popup.isModal())
        self.assertEqual(self.window.stack.currentIndex(), current)
        self.assertEqual(self.window.nav_group.checkedId(), 0)
        from PySide6.QtWidgets import QCheckBox
        hold = popup.findChild(QCheckBox)
        hold.setChecked(False)
        self.window.show_page(5)
        self.assertTrue(popup.isVisible())
        popup.close()
        self.window.show_page(3)
        self.assertIs(self.window.spectrum_window, popup)
        self.assertFalse(hold.isChecked())
        self.window.close()
        self.assertFalse(popup.isVisible())

    def test_dedicated_log_page_preserves_selected_details_as_messages_arrive(self):
        self.connect_first()
        self.window.preview_send()
        self.assertFalse(self.window.stack.currentWidget().isAncestorOf(self.window.log_table))
        self.window.show_page(5)
        self.assertTrue(self.window.stack.currentWidget().isAncestorOf(self.window.log_table))
        self.assertEqual(self.window.log_table.rowCount(), 1)
        self.window.log_table.selectRow(0)
        selected_id = self.window.log_events[0]['event_id']
        self.window.mock_receive()
        self.window.show_log_detail()  # Arrivals must not shift details before the batched redraw.
        self.assertEqual(json.loads(self.window.log_detail.toPlainText())['event_id'], selected_id)
        self.window.flush_log_view()
        self.assertEqual(self.window.log_table.currentRow(), 1)
        self.assertEqual(json.loads(self.window.log_detail.toPlainText())['event_id'], selected_id)

    def test_monitor_pause_freezes_display_but_log_collection_continues(self):
        self.connect_first()
        self.window.preview_send()
        monitor = self.window.traffic_monitor
        monitor.render_now()
        QTest.mouseClick(monitor.pause_button, Qt.MouseButton.LeftButton)
        snapshot = monitor.paused_snapshot
        self.window.mock_receive()
        monitor.render_now()
        self.assertEqual(len(monitor.canvas.records), 1)
        self.assertEqual(snapshot.stats()['RX'], 0)
        self.assertEqual(len(self.window.log_events), 2)
        self.window.flush_log_view()
        self.assertEqual(self.window.log_table.rowCount(), 2)
        QTest.mouseClick(monitor.pause_button, Qt.MouseButton.LeftButton)
        self.assertIsNone(monitor.paused_snapshot)
        self.assertEqual(len(monitor.canvas.records), 2)

    def test_timeline_click_exposes_actual_message_values(self):
        self.connect_first()
        self.window.payload_inputs['channel'].setText('3')
        self.window.preview_send()
        monitor = self.window.traffic_monitor
        monitor.render_now()
        monitor.canvas.grab()  # Force a layout/paint even if the parent scroll area clips it.
        self.assertTrue(monitor.canvas.hits)
        rect, records = monitor.canvas.hits[0]
        QTest.mouseClick(monitor.canvas, Qt.MouseButton.LeftButton, pos=rect.center().toPoint())
        shown = json.loads(monitor.detail.toPlainText())
        self.assertEqual(shown['events'][0]['values']['channel'], 3)
        self.assertEqual(shown['events'][0]['mode'], 'MOCK')
        self.assertFalse(any(key.startswith('_') for key in shown['events'][0]))

    def test_monitor_filter_zoom_and_history_controls(self):
        self.connect_first()
        self.window.preview_send()
        self.window.select_connection(1)
        self.window.message_catalog.setCurrentRow(0)
        self.window.preview_send()
        monitor = self.window.traffic_monitor
        monitor.connection_filter.setCurrentIndex(monitor.connection_filter.findData('TCP-01'))
        self.assertEqual(monitor.canvas.lanes, [('TCP-01', '')])
        monitor.span.setCurrentIndex(0)
        self.assertAlmostEqual(monitor.canvas.end - monitor.canvas.start, 5)
        monitor.follow.setChecked(False)
        monitor.offset.setValue(120)
        self.assertEqual(monitor.canvas.lanes, [])
        monitor.follow.setChecked(True)
        self.assertEqual(monitor.canvas.lanes, [('TCP-01', '')])

    def test_log_buffer_overflow_keeps_counters_and_clears_evicted_selection(self):
        self.connect_first()
        self.window.preview_send()
        self.window.show_page(5)
        self.window.log_table.selectRow(0)
        original = dict(self.window.log_events[0])
        for _ in range(2001):
            self.window.record_event(original)
        self.window.flush_log_view()
        self.assertEqual(self.window.log_table.rowCount(), 2000)
        self.assertFalse(self.window.log_table.selectionModel().hasSelection())
        self.assertEqual(self.window.traffic_history.snapshot().stats()['TX'], 2002)

    def choose_excel_sheet(self, path, index=0, cancel=False):
        from ds_simulator.icd_ui import SheetSelectionDialog
        from PySide6.QtWidgets import QDialogButtonBox
        def choose():
            dialog = next(w for w in APP.topLevelWidgets() if isinstance(w, SheetSelectionDialog) and w.isVisible())
            dialog.sheets.setCurrentRow(index)
            action = QDialogButtonBox.StandardButton.Cancel if cancel else QDialogButtonBox.StandardButton.Ok
            QTest.mouseClick(dialog.actions.button(action), Qt.MouseButton.LeftButton)
        QTimer.singleShot(30, choose)
        with patch('ds_simulator.icd_ui.QFileDialog.getOpenFileName', return_value=(str(path), 'Excel (*.xlsx *.xlsm)')):
            self.window.import_excel_dialog()

    def test_excel_chooser_imports_only_selected_sheet_and_composer_uses_choices(self):
        from test_icd import make_workbook
        from PySide6.QtWidgets import QComboBox, QLineEdit
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'synthetic.xlsx'; make_workbook(path)
            self.choose_excel_sheet(path, index=1)
        self.assertEqual(len(self.window.messages), 1)
        self.assertEqual(self.window.message_catalog.item(0).text(), '0x0020\n가상 출력 설정')
        self.assertEqual(self.window.scenario_editor.tree.topLevelItemCount(), 0)
        self.window.message_catalog.setCurrentRow(0)  # Inspect without a connection.
        choice = self.window.payload_inputs['field_0001']
        free = self.window.payload_inputs['field_0002']
        self.assertIsInstance(choice, QComboBox)
        self.assertIsInstance(free, QLineEdit)
        self.assertTrue(free.isReadOnly())
        self.assertEqual([choice.itemText(i) for i in range(choice.count())], ['LOW', 'MEDIUM', 'HIGH'])
        self.assertFalse(self.window.send_button.isEnabled())
        choice.showPopup(); APP.processEvents()
        index = choice.model().index(2, 0)
        QTest.mouseClick(choice.view().viewport(), Qt.MouseButton.LeftButton, pos=choice.view().visualRect(index).center())
        self.assertEqual(choice.currentText(), 'HIGH')
        self.window.toggle_preview_connection()
        self.assertTrue(self.window.send_button.isEnabled())
        QTest.mouseClick(self.window.send_button, Qt.MouseButton.LeftButton)
        record = self.window.log_events[0]
        self.assertEqual(record['message'], '가상 출력 설정')
        self.assertEqual(record['headers']['cmd'], '0x0020')
        self.assertEqual(record['values']['field_0001'], 'HIGH')
        self.assertEqual(record['values']['field_0002'], '')
        self.assertEqual(record['data_elements'][0]['fixed']['E'], '고정 0 1')
        self.assertEqual(record['data_elements'][0]['fixed']['K'], '원본 유지')
        self.assertEqual(record['data_elements'][0]['name'], '출력 모드')
        self.assertEqual(self.window.definition_catalog.rowCount(), 1)
        self.assertEqual(self.window.definition_fields.rowCount(), 2)

    def test_excel_cancel_or_parse_error_preserves_previous_catalog_and_scenario(self):
        from test_icd import make_workbook
        from openpyxl import load_workbook
        original = self.window.scenario_editor.collect().to_dict()
        before = dict(self.window.messages)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'synthetic.xlsx'; make_workbook(path)
            self.choose_excel_sheet(path, cancel=True)
            self.assertEqual(self.window.messages, before)
            book = load_workbook(path); book['가상 장치 A(데이터)']['E2'] = '=1+1'; book.save(path); book.close()
            with patch('ds_simulator.icd_ui.QMessageBox.warning') as warning:
                self.choose_excel_sheet(path)
                warning.assert_called_once()
        self.assertEqual(self.window.messages, before)
        self.assertEqual(self.window.scenario_editor.collect().to_dict(), original)

    def test_fixed_icd_table_keeps_c_d_and_m_visible_and_scrolls_rows_together(self):
        from test_icd import HEADERS, icd_row
        from openpyxl import Workbook
        from ds_simulator.icd import read_icd_sheet
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'table.xlsx'
            book = Workbook(); sheet = book.active; sheet.title = '표(데이터)'
            sheet.append(HEADERS)
            for i in range(30):
                sheet.append(icd_row('0x1' if i == 0 else None, '가상 표' if i == 0 else None, i + 1, f'요소 {i + 1}', '선택 A\n선택 B'))
            book.save(path); book.close()
            catalog = read_icd_sheet(path, sheet.title)
            self.window.apply_catalog(catalog, '가상 표')
        self.window.message_catalog.setCurrentRow(0)
        table = self.window.payload_table
        table.setFixedHeight(240); APP.processEvents()
        for view in (table.identity, table.details):
            for column in range(view.columnCount()):
                self.assertFalse(view.item(0, column).flags() & Qt.ItemFlag.ItemIsEditable)
        before = table.details.item(0, 0).text()
        table.details.setFocus()
        QTest.keyClicks(table.details, 'blocked')
        self.assertEqual(table.details.item(0, 0).text(), before)
        left = table.identity.geometry(); right = table.settings.geometry()
        table.details.horizontalScrollBar().setValue(table.details.horizontalScrollBar().maximum())
        self.assertEqual(table.identity.geometry(), left)
        self.assertEqual(table.settings.geometry(), right)
        table.details.verticalScrollBar().setValue(150)
        self.assertEqual(table.identity.verticalScrollBar().value(), 150)
        self.assertEqual(table.settings.verticalScrollBar().value(), 150)
        self.assertEqual(table.settings.cellWidget(0, 0).count(), 2)
        self.assertIn('고정 설명', self.window.definition_fields.horizontalHeaderItem(2).text())
        self.assertEqual(self.window.definition_fields.item(0, 9).text(), '선택 A\n선택 B')

    def test_m_line_labels_keep_spaces_and_blank_m_cannot_be_typed_into(self):
        from test_icd import HEADERS, icd_row
        from openpyxl import Workbook
        from ds_simulator.icd import read_icd_sheet
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'm.xlsx'
            book = Workbook(); sheet = book.active; sheet.title = '선택(데이터)'
            sheet.append(HEADERS)
            sheet.append(icd_row('0x1', '가상 선택', 1, '운용', '저속 운용\n고속  운용'))
            sheet.append(icd_row(None, None, 2, '고정 항목', None))
            book.save(path); book.close()
            catalog = read_icd_sheet(path, sheet.title)
        self.window.apply_catalog(catalog, '가상 선택')
        self.window.message_catalog.setCurrentRow(0)
        choice = self.window.payload_inputs['field_0001']
        empty = self.window.payload_inputs['field_0002']
        self.assertEqual([choice.itemText(i) for i in range(choice.count())], ['저속 운용', '고속  운용'])
        self.assertFalse(choice.isEditable())
        empty.setFocus(); QTest.keyClicks(empty, 'blocked')
        self.assertEqual(empty.text(), '')
        choice.setCurrentText('고속  운용')
        self.window.toggle_preview_connection()
        self.window.preview_send()
        event = self.window.log_events[0]
        self.assertEqual(event['values']['field_0001'], '고속  운용')
        self.assertEqual(event['data_elements'][0]['fixed']['E'], '고정 0 1')
        self.assertEqual(len(event['data_elements'][0]['fixed']), 9)

    def test_excel_formula_opcode_uses_saved_result_in_list_and_mock_send(self):
        from openpyxl import Workbook
        from test_icd import set_formula_caches, HEADERS, icd_row
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'formula.xlsx'
            book = Workbook(); sheet = book.active; sheet.title = '수식(데이터)'
            sheet.append(HEADERS)
            sheet.append(icd_row('="0x00A1"', '가상 수식 설정', 1, '운용 모드', 'OFF\nON\nAUTO'))
            book.save(path); book.close()
            set_formula_caches(path, {'A2': ('str', '0x00A1')})
            self.choose_excel_sheet(path)
        self.assertEqual(self.window.message_catalog.item(0).text(), '0x00A1\n가상 수식 설정')
        self.window.message_catalog.setCurrentRow(0)
        self.window.payload_inputs['field_0001'].setCurrentText('ON')
        self.window.toggle_preview_connection()
        QTest.mouseClick(self.window.send_button, Qt.MouseButton.LeftButton)
        record = self.window.log_events[0]
        self.assertEqual(record['headers']['cmd'], '0x00A1')
        self.assertEqual(record['values']['field_0001'], 'ON')

    def test_imported_scenario_restores_choices_and_runs_without_excel(self):
        from test_icd import make_workbook
        from ds_simulator.icd import read_icd_sheet
        from ds_simulator.preview_data import MESSAGES
        from PySide6.QtWidgets import QComboBox
        self.connect_first()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'synthetic.xlsx'; make_workbook(path)
            catalog = read_icd_sheet(path, '가상 장치 B(데이터)')
            self.window.apply_catalog(catalog, '가상 선택 시트')
            editor = self.window.scenario_editor
            editor.palette.setCurrentRow(2); editor.add_selected()
            from ds_simulator.icd_widgets import IcdFieldTable
            from PySide6.QtWidgets import QDialog
            step = editor.tree.currentItem().data(0, ROLE)
            readonly = []
            def choose():
                dialog = next(w for w in APP.topLevelWidgets() if isinstance(w, QDialog) and w.isVisible())
                try:
                    table = dialog.findChild(IcdFieldTable)
                    table.inputs['field_0001'].setCurrentText('HIGH')
                    readonly.append(table.inputs['field_0002'].isReadOnly())
                finally:
                    dialog.accept()
            QTimer.singleShot(30, choose)
            editor.edit_icd_table(step, catalog[step.message])
            self.assertEqual(readonly, [True])
            output = Path(directory) / 'scenario.json'
            editor.save_to(output)
            with patch('sys.excepthook') as exceptions:
                self.window.apply_catalog(MESSAGES, '기본 예시')
                path.unlink()
                editor.load_from(output)
                self.assertEqual(self.window.message_catalog.item(0).text(), '0x0020\n가상 출력 설정')
                editor.run(); wait_until(lambda: not editor.runner.running)
                self.assertEqual(self.window.log_events[0]['values']['field_0001'], 'HIGH')
                self.assertEqual(editor.runner.completed, 1)
                self.assertEqual(editor.collect().to_dict()['version'], 2)
                APP.processEvents()
                exceptions.assert_not_called()

    def test_imported_catalog_is_scoped_to_one_window(self):
        from test_icd import make_workbook
        from ds_simulator.icd import read_icd_sheet
        from ds_simulator.preview_data import MESSAGES
        other = SimulatorWindow(mock=True)
        try:
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / 'synthetic.xlsx'; make_workbook(path)
                self.window.apply_catalog(read_icd_sheet(path, '가상 장치 B(데이터)'), '가상 시트')
            self.assertEqual(len(self.window.messages), 1)
            self.assertEqual(len(other.messages), 5)
            self.assertEqual(len(MESSAGES), 5)
        finally:
            other.scenario_editor.dirty = False; other.close(); other.deleteLater()

    def test_excel_import_is_blocked_while_scenario_is_running(self):
        self.connect_first()
        editor = self.window.scenario_editor
        editor.document = Scenario(steps=[Step('wait', milliseconds=1000)]); editor.populate(); editor.run()
        with patch('ds_simulator.icd_ui.QFileDialog.getOpenFileName') as picker, patch('ds_simulator.icd_ui.QMessageBox.information'):
            self.window.import_excel_dialog()
            picker.assert_not_called()
        editor.runner.stop()

    def test_excel_with_no_data_sheet_leaves_existing_ui_intact(self):
        from openpyxl import Workbook
        before = dict(self.window.messages)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'no-data.xlsx'; book = Workbook(); book.save(path); book.close()
            with patch('ds_simulator.icd_ui.QFileDialog.getOpenFileName', return_value=(str(path), 'Excel')), patch('ds_simulator.icd_ui.QMessageBox.warning') as warning:
                self.window.import_excel_dialog()
                warning.assert_called_once()
        self.assertEqual(self.window.messages, before)

    def test_manual_send_validation_and_actual_values_in_log(self):
        self.connect_first()
        self.window.payload_inputs['channel'].setText('256')
        self.assertFalse(self.window.send_button.isEnabled())
        self.window.payload_inputs['channel'].setText('2')
        self.window.big_endian.setChecked(True)
        QTest.mouseClick(self.window.send_button, Qt.MouseButton.LeftButton)
        self.assertEqual(len(self.window.log_events), 1)
        event = self.window.log_events[0]
        self.assertEqual(event['values']['channel'], 2)
        self.assertTrue(event['big_endian'])
        self.assertEqual(event['connection_id'], 'TCP-01')
        self.assertEqual(event['source'], '수동')

    def test_add_move_indent_outdent_and_save_reload(self):
        editor = self.window.scenario_editor
        root = editor.tree.topLevelItem(0)
        editor.tree.setCurrentItem(root)
        editor.palette.setCurrentRow(1)  # wait
        editor.add_selected()
        self.assertEqual(root.childCount(), 4)
        added = editor.tree.currentItem()
        editor.move(-1)
        self.assertIs(root.child(2), added)
        editor.outdent()
        self.assertIs(editor.tree.topLevelItem(1), added)
        editor.indent()
        self.assertIs(added.parent(), root)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / '한글.json'
            editor.save_to(path)
            saved = editor.collect().to_dict()
            editor.document = Scenario('empty')
            editor.populate()
            editor.load_from(path)
            self.assertEqual(editor.collect().to_dict(), saved)
            self.assertFalse(editor.dirty)

    def test_invalid_load_preserves_editor(self):
        editor = self.window.scenario_editor
        original = editor.collect().to_dict()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'broken.json'
            path.write_text('{broken')
            with self.assertRaises(ValueError):
                editor.load_from(path)
        self.assertEqual(editor.collect().to_dict(), original)

    def test_scenario_execution_reaches_shared_log_and_locks_editor(self):
        self.connect_first()
        self.window.show_page(4)
        editor = self.window.scenario_editor
        repeat = editor.tree.topLevelItem(0).data(0, ROLE)
        editor.tree.setCurrentItem(editor.tree.topLevelItem(0).child(0))
        from PySide6.QtWidgets import QComboBox
        target = editor.properties.findChild(QComboBox)
        self.assertNotIn('재선택', target.currentText())
        repeat.count = 2
        editor.tree.topLevelItem(0).child(1).data(0, ROLE).milliseconds = 1
        QTest.mouseClick(editor.run_button, Qt.MouseButton.LeftButton)
        self.assertFalse(editor.edit_area.isEnabled())
        wait_until(lambda: not editor.runner.running)
        self.assertTrue(editor.edit_area.isEnabled())
        self.assertEqual(len(self.window.log_events), 4)
        self.assertEqual(editor.runner.completed, 6)
        self.assertTrue(all(e['source'] == '시나리오' for e in self.window.log_events))

    def test_profiles_roundtrip_clears_sessions_and_requires_reconnect(self):
        self.connect_first()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'connections.json'
            self.window.save_profiles_to(path)
            expected = list(self.window.connections)
            self.window.load_profiles_from(path)
        self.assertEqual(self.window.connections, expected)
        self.assertEqual(self.window.sockets, [])
        self.assertFalse(self.window.send_button.isEnabled())
        self.assertEqual(self.window.socket_table.rowCount(), 0)

    def test_filtered_export_contains_values_and_mock_marker(self):
        self.connect_first()
        self.window.preview_send()
        self.window.mock_receive()
        self.assertEqual(len(self.window.log_events), 2)
        self.window.log_direction.setCurrentIndex(1)  # TX only
        self.assertEqual(len(self.window.visible_events()), 1)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'log.json'
            with patch('ds_simulator.activity.QFileDialog.getSaveFileName', return_value=(str(path), 'JSON (*.json)')):
                self.window.export_log()
            data = json.loads(path.read_text())
            self.assertEqual(len(data['events']), 1)
            self.assertEqual(data['events'][0]['mode'], 'MOCK')
            self.assertEqual(data['events'][0]['values']['center_frequency'], 1000000000)
            path = Path(directory) / 'log.csv'
            with patch('ds_simulator.activity.QFileDialog.getSaveFileName', return_value=(str(path), 'CSV (*.csv)')):
                self.window.export_log()
            with path.open(encoding='utf-8-sig', newline='') as stream:
                rows = list(csv.reader(stream))
            self.assertNotIn(b'\r\r\n', path.read_bytes())
            self.assertEqual(len(rows), 2)
            self.assertIn('MOCK', rows[1][7])
