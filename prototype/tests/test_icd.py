"""All workbooks in this suite are synthetic and contain no customer ICD data."""
from datetime import datetime
import hashlib
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from ds_simulator.icd import IcdMessage, data_sheet_names, read_icd_sheet, catalog_payload, catalog_from_payload
from ds_simulator.domain import Scenario, Step, load_scenario, save_json, validate_message
from ds_simulator.preview_data import MESSAGES


def make_workbook(path):
    book = Workbook()
    notes = book.active
    notes.title = '설명'
    notes.append(['이 시트는 ICD가 아닙니다'])
    notes.append(['IMPORT-MUST-NOT-SEE-THIS', '설명 내용', 1, '무시할 필드', 'X Y'])
    first = book.create_sheet('가상 장치 A(데이터)')
    first.append(['opcode', '연동자료명', '순번', '데이터 요소', '데이터 상세 값'])
    first.append(['0x0010', '가상 운용 설정', 1, '운용 모드', 'OFF ON\tAUTO\nHOLD'])
    first.append([None, None, 2, '채널', '1 2 3'])
    first.merge_cells('A2:A3'); first.merge_cells('B2:B3')
    first.append(['0x0010', '가상 운용 설정', 3, '비고', None])
    first.append(['0x0011', '가상 상태 조회', 1, '조회 대상', 'A B A'])
    second = book.create_sheet('가상 장치 B(데이터)')
    second.append(['opcode', '연동자료명', '순번', '데이터 요소', '데이터 상세 값'])
    second.append(['0x0020', '가상 출력 설정', 1, '출력 모드', 'LOW MEDIUM HIGH'])
    second.append([None, None, 2, '설명', None])
    book.create_sheet('이전(데이터) 백업').append(['불러오지 않습니다'])
    book.save(path)
    book.close()


class IcdReaderTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.path = Path(self.directory.name) / 'synthetic-icd.xlsx'
        make_workbook(self.path)

    def tearDown(self):
        self.directory.cleanup()

    def test_only_data_suffix_sheets_are_listed_and_only_chosen_sheet_is_imported(self):
        self.assertEqual(data_sheet_names(self.path), ['가상 장치 A(데이터)', '가상 장치 B(데이터)'])
        messages = read_icd_sheet(self.path, '가상 장치 B(데이터)')
        self.assertEqual([(m.opcode, m.description) for m in messages.values()], [('0x0020', '가상 출력 설정')])
        with self.assertRaises(ValueError):
            read_icd_sheet(self.path, '설명')

    def test_merged_and_repeated_identity_rows_are_grouped_without_losing_fields(self):
        messages = read_icd_sheet(self.path, '가상 장치 A(데이터)')
        self.assertEqual(len(messages), 2)
        message = next(iter(messages.values()))
        self.assertEqual([f.sequence for f in message.elements], ['1', '2', '3'])
        self.assertEqual([f.name for f in message.elements], ['운용 모드', '채널', '비고'])
        self.assertEqual(message.elements[0].choices, ('OFF', 'ON', 'AUTO', 'HOLD'))
        self.assertEqual(message.elements[2].choices, ())
        self.assertEqual(list(messages.values())[1].elements[0].choices, ('A', 'B'))

    def test_numeric_opcode_zero_and_excel_zero_padding_are_preserved(self):
        book = Workbook(); sheet = book.active; sheet.title = '숫자(데이터)'
        sheet.append(['header'] * 5); sheet.append([0, '가상 정지', 1, '선택', '0 1'])
        sheet['A2'].number_format = '0000'
        book.save(self.path); book.close()
        message = next(iter(read_icd_sheet(self.path, sheet.title).values()))
        self.assertEqual(message.opcode, '0000')
        self.assertEqual(message.elements[0].choices, ('0', '1'))

    def test_new_opcode_does_not_silently_inherit_previous_message_name(self):
        book = Workbook(); sheet = book.active; sheet.title = '가상(데이터)'
        sheet.append(['header'] * 5)
        sheet.append(['0x1', '가상 첫 메시지', 1, '요소', 'A B'])
        sheet.append(['0x2', None, 1, '다른 요소', 'C D'])
        book.save(self.path); book.close()
        with self.assertRaisesRegex(ValueError, '3행'):
            read_icd_sheet(self.path, sheet.title)

    def test_only_a_to_e_are_read_and_formulas_are_not_executed(self):
        book = Workbook(); sheet = book.active; sheet.title = '가상(데이터)'
        sheet.append(['header'] * 6)
        sheet.append(['0x1', '가상 메시지', 1, '요소', 'A B', '=1+1'])
        book.save(self.path)
        self.assertEqual(len(read_icd_sheet(self.path, sheet.title)), 1)
        sheet['E2'] = '=1+1'; book.save(self.path); book.close()
        with self.assertRaisesRegex(ValueError, 'E2'):
            read_icd_sheet(self.path, sheet.title)

    def test_duplicate_element_names_have_independent_values_and_file_is_unchanged(self):
        book = Workbook(); sheet = book.active; sheet.title = '가상(데이터)'
        sheet.append(['header'] * 5)
        sheet.append(['0x1', '가상 메시지', 1, '같은 이름', 'A B'])
        sheet.append([None, None, 2, '같은 이름', 'C D'])
        book.save(self.path); book.close()
        before = hashlib.sha256(self.path.read_bytes()).hexdigest()
        message = next(iter(read_icd_sheet(self.path, sheet.title).values()))
        self.assertNotEqual(message.elements[0].key, message.elements[1].key)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(), before)

    def test_catalog_and_scenario_roundtrip_does_not_require_the_original_workbook(self):
        catalog = read_icd_sheet(self.path, '가상 장치 B(데이터)')
        key = next(iter(catalog)); step = Step.send(key, catalog)
        step.values['field_0001'] = 'HIGH'; step.values['field_0002'] = '직접 입력한 값'
        scenario = Scenario('가상 엑셀 시험', [step], catalog=catalog)
        path = self.path.with_suffix('.json')
        save_json(path, scenario.to_dict()); self.path.unlink()
        restored = load_scenario(path)
        self.assertEqual(restored.to_dict(), scenario.to_dict())
        self.assertEqual(restored.to_dict()['version'], 2)
        self.assertEqual(catalog_from_payload(catalog_payload(catalog)), catalog)

    def test_choice_validation_rejects_undefined_values_without_guessing_wire_types(self):
        catalog = read_icd_sheet(self.path, '가상 장치 B(데이터)')
        key = next(iter(catalog)); step = Step.send(key, catalog)
        step.values['field_0001'] = 'HIGH'
        self.assertEqual(validate_message(key, step.headers, step.values, catalog)['field_0001'], 'HIGH')
        step.values['field_0001'] = 'NOT_AN_OPTION'
        with self.assertRaises(ValueError):
            validate_message(key, step.headers, step.values, catalog)

    def test_invalid_excel_is_a_readable_error(self):
        self.path.write_text('not a zip workbook')
        with self.assertRaises(ValueError):
            data_sheet_names(self.path)
